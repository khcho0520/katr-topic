"""논문(PDF/DOCX/TXT) → 정제 → 청킹 → Corpus.

분석 단위 = 청크. 두 데이터 체제(헤드라인/논문)를 모델 입장에서 구조적으로
동일하게 만들어 공정 비교를 가능하게 한다. 논문별 집계는 보고용 뷰에서 수행.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import List, Tuple, Union

from core.data.cleaning import clean_document_text
from core.types import Corpus, Document
from core.utils.hashing import texts_hash

DEFAULT_CHUNK_CHARS = 900
DEFAULT_OVERLAP_CHARS = 100


def extract_text(source: Union[str, Path, bytes], filename: str) -> str:
    """PDF/DOCX/TXT에서 원문 텍스트 추출."""
    name = filename.lower()
    data: bytes
    if isinstance(source, (str, Path)):
        data = Path(source).read_bytes()
    else:
        data = source

    if name.endswith(".pdf"):
        import fitz  # pymupdf

        with fitz.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text("text") for page in doc)
    if name.endswith(".docx"):
        from docx import Document as DocxDocument

        docx = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in docx.paragraphs)
    # txt/md 등
    for encoding in ("utf-8", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def chunk_text(
    text: str,
    chunk_chars: int = DEFAULT_CHUNK_CHARS,
    overlap_chars: int = DEFAULT_OVERLAP_CHARS,
) -> List[str]:
    """문단 경계를 우선 존중하며 chunk_chars 내외로 분할, 청크 간 overlap 유지."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: List[str] = []
    current = ""
    for para in paragraphs:
        if current and len(current) + len(para) + 1 > chunk_chars:
            chunks.append(current.strip())
            current = current[-overlap_chars:] if overlap_chars > 0 else ""
        if len(para) > chunk_chars * 1.5:
            # 매우 긴 문단은 문장 단위로 강제 분할
            sentences = re.split(r"(?<=[.!?다\.])\s+", para)
            for sent in sentences:
                if current and len(current) + len(sent) + 1 > chunk_chars:
                    chunks.append(current.strip())
                    current = current[-overlap_chars:] if overlap_chars > 0 else ""
                current = f"{current} {sent}".strip()
        else:
            current = f"{current}\n{para}".strip() if current else para
    if current.strip():
        chunks.append(current.strip())
    return [c for c in chunks if len(c) >= 50]


def load_papers(
    files: List[Tuple[Union[str, Path, bytes], str]],
    chunk_chars: int = DEFAULT_CHUNK_CHARS,
    overlap_chars: int = DEFAULT_OVERLAP_CHARS,
    clean: bool = True,
    name: str = "",
) -> Corpus:
    """files: [(파일 경로 또는 bytes, 파일명), ...] → 청크 단위 Corpus."""
    docs: List[Document] = []
    per_file_chunks: dict[str, int] = {}
    for source, filename in files:
        raw = extract_text(source, filename)
        cleaned = clean_document_text(raw) if clean else raw
        chunks = chunk_text(cleaned, chunk_chars, overlap_chars)
        source_id = Path(filename).stem
        per_file_chunks[source_id] = len(chunks)
        for ci, chunk in enumerate(chunks):
            docs.append(
                Document(
                    doc_id=f"{source_id}#c{ci}",
                    text=chunk,
                    source_id=source_id,
                    meta={"chunk_index": ci},
                )
            )

    if not docs:
        raise ValueError("추출된 청크가 없습니다 (파일 내용/청킹 설정 확인).")

    return Corpus(
        regime="papers",
        docs=docs,
        # 청킹 파라미터를 해시에 포함 → 설정 변경 시 캐시가 올바르게 무효화됨
        content_hash=texts_hash([d.text for d in docs] + [f"chunk:{chunk_chars}:{overlap_chars}"]),
        name=name or "papers",
        meta={
            "n_files": len(files),
            "chunk_chars": chunk_chars,
            "overlap_chars": overlap_chars,
            "clean": clean,
            "per_file_chunks": per_file_chunks,
        },
    )
